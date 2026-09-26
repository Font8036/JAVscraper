"""按列名读取 Excel。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from .config import JellyfinNfoConfig
from .parser import (
    parse_actors, parse_genres, parse_rating, split_columns,
)


def _build_header_map(ws, header_row: int) -> dict:
    headers: dict = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=col).value
        if v is not None:
            headers[str(v).strip()] = col
    return headers


def _resolve_cols(headers: dict, name_str: str) -> list:
    return [
        headers[n] for n in split_columns(name_str) if n in headers
    ]


def read_excel_data(config: JellyfinNfoConfig) -> dict:
    """读取 Excel，返回 {code: movie_data}。"""
    path = Path(config.input_excel)
    if not path.exists():
        raise FileNotFoundError(f"Excel 不存在：{path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if config.sheet_index < 1 or config.sheet_index > len(wb.worksheets):
            raise ValueError(
                f"工作表索引 {config.sheet_index} 超出范围"
                f"（共 {len(wb.worksheets)} 个）"
            )
        ws = wb.worksheets[config.sheet_index - 1]
        headers = _build_header_map(ws, config.header_row)

        code_col = headers.get(config.code_column)
        if code_col is None:
            raise ValueError(
                f"找不到番号列：{config.code_column}\n"
                f"当前表头行可用列：{list(headers)}"
            )

        title_col = headers.get(config.title_column)
        actor_cols = _resolve_cols(headers, config.actor_columns)
        rating_cols = _resolve_cols(headers, config.rating_columns)
        genre_cols = _resolve_cols(headers, config.genre_columns)
        personal_cols = _resolve_cols(headers, config.personal_comment_columns)
        user_cols = _resolve_cols(headers, config.user_comment_columns)

        result: dict = {}

        for row in range(config.header_row + 1, ws.max_row + 1):
            code_val = ws.cell(row=row, column=code_col).value
            if code_val is None:
                continue
            code = str(code_val).strip()
            if not code:
                continue

            # ---- 标题 ----
            title = code
            if title_col is not None:
                tv = ws.cell(row=row, column=title_col).value
                if tv is not None and str(tv).strip():
                    title = str(tv).strip()

            # ---- 演员：所有列合并，按 name 去重 ----
            actors: list = []
            seen_names: set = set()
            for c in actor_cols:
                v = ws.cell(row=row, column=c).value
                if v is None:
                    continue
                for a in parse_actors(v):
                    if a["name"] and a["name"] not in seen_names:
                        seen_names.add(a["name"])
                        actors.append(a)

            # ---- 类别：所有列取并集 ----
            genre_set: set = set()
            for c in genre_cols:
                v = ws.cell(row=row, column=c).value
                if v is None:
                    continue
                genre_set.update(parse_genres(v))
            genres = sorted(genre_set)

            # ---- 评分：所有列按顺序保留 ----
            ratings: list = []
            for c in rating_cols:
                v = ws.cell(row=row, column=c).value
                if v is None:
                    continue
                score, count = parse_rating(v)
                if score is not None:
                    ratings.append((score, count))

            # ---- 个人评论：按列顺序前后相接 ----
            personal_parts: list = []
            for c in personal_cols:
                v = ws.cell(row=row, column=c).value
                if v is not None and str(v).strip():
                    personal_parts.append(str(v).strip())
            personal_comment = "\n".join(personal_parts)

            # ---- 网友评论：按列顺序前后相接 ----
            user_parts: list = []
            for c in user_cols:
                v = ws.cell(row=row, column=c).value
                if v is not None and str(v).strip():
                    user_parts.append(str(v).strip())
            user_comment = "\n".join(user_parts)

            result[code] = {
                "code": code,
                "title": title,
                "actors": actors,
                "ratings": ratings,
                "genres": genres,
                "personal_comment": personal_comment,
                "user_comment": user_comment,
            }

        return result
    finally:
        wb.close()