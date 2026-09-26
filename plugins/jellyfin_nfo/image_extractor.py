"""从 Excel 中自动提取封面（支持 WPS 的 DISPIMG 嵌入方式）。

不需要用户手动解压 xlsx —— 内部用 zipfile 直接读取。
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Callable, Optional

import openpyxl

LogFn = Callable[[str], None]

_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def extract_covers_from_excel(
    excel_path: str | Path,
    sheet_index: int,
    code_column_name: str,
    cover_column_name: str,          # ← 新增
    header_row: int,
    output_dir: str | Path,
    on_log: Optional[LogFn] = None,
) -> tuple:
    """从 Excel 提取封面，按番号命名保存到 output_dir。

    返回 (导出数, 跳过数)。
    """
    log = on_log or (lambda s: None)
    excel_path = Path(excel_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    # 1. 番号 → DISPIMG_ID
    code_to_dispimg = _read_dispimg_map(
        excel_path, sheet_index, code_column_name, cover_column_name, header_row,
    )
    log(f"从表格读取到 {len(code_to_dispimg)} 个 DISPIMG 引用")

    if not code_to_dispimg:
        log("未找到任何 DISPIMG 图片。")
        log("提示：本功能要求表格使用了「嵌入单元格图片」的保存方式。")
        log("如果用 WPS 编辑过表格，请保存为 .xlsx 后再试。")
        return 0, 0

    # 2. DISPIMG_ID → 图片文件名
    dispimg_to_file = _parse_cellimages(excel_path)
    log(f"从表格内部建立 {len(dispimg_to_file)} 个 ID → 图片文件的映射")

    # 3. 提取图片
    exported = 0
    skipped = 0

    with zipfile.ZipFile(excel_path) as zf:
        names = set(zf.namelist())
        for code, dispimg_id in code_to_dispimg.items():
            filename = dispimg_to_file.get(dispimg_id)
            if not filename:
                log(f"跳过 {code}: 无对应图片")
                skipped += 1
                continue

            inner = f"xl/media/{filename}"
            if inner not in names:
                log(f"跳过 {code}: 图片 {inner} 不存在")
                skipped += 1
                continue

            ext = Path(filename).suffix or ".jpg"
            target = output_dir / f"{code}{ext}"
            try:
                target.write_bytes(zf.read(inner))
                log(f"已导出 {code}{ext}")
                exported += 1
            except OSError as e:
                log(f"写入 {target} 失败: {e}")
                skipped += 1

    return exported, skipped


# ============================================================
# 内部
# ============================================================
def _read_dispimg_map(
    excel_path: Path, sheet_index: int,
    code_column_name: str, cover_column_name: str, header_row: int,
) -> dict:
    """{番号: DISPIMG_ID}"""
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    try:
        if sheet_index < 1 or sheet_index > len(wb.worksheets):
            raise ValueError(f"工作表索引 {sheet_index} 超出范围")
        ws = wb.worksheets[sheet_index - 1]

        # 定位列
        header_to_col: dict = {}
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=col).value
            if v is not None:
                header_to_col[str(v).strip()] = col

        code_col = header_to_col.get(code_column_name)
        if code_col is None:
            raise ValueError(
                f"找不到番号列：{code_column_name}（在表头行 {header_row}）\n"
                f"当前表头可用列：{list(header_to_col)}"
            )

        cover_col = header_to_col.get(cover_column_name)
        if cover_col is None:
            raise ValueError(
                f"找不到封面列：{cover_column_name}（在表头行 {header_row}）\n"
                f"当前表头可用列：{list(header_to_col)}"
            )

        result: dict = {}
        for row in range(header_row + 1, ws.max_row + 1):
            code_val = ws.cell(row=row, column=code_col).value
            if code_val is None:
                continue
            code = str(code_val).strip()
            if not code:
                continue

            # 只看封面列
            cell = ws.cell(row=row, column=cover_col)
            if cell.data_type == "f" and isinstance(cell.value, str):
                if "DISPIMG" in cell.value:
                    m = re.search(r'DISPIMG\("([^"]+)"', cell.value)
                    if m:
                        result[code] = m.group(1)

        return result
    finally:
        wb.close()


def _parse_cellimages(excel_path: Path) -> dict:
    """{DISPIMG_ID: 图片文件名}"""
    with zipfile.ZipFile(excel_path) as zf:
        names = set(zf.namelist())

        cellimages_path = "xl/cellimages.xml"
        rels_path = "xl/_rels/cellimages.xml.rels"
        if cellimages_path not in names:
            return {}

        # DISPIMG_ID → rId
        dispimg_to_rid: dict = {}
        with zf.open(cellimages_path) as f:
            root = ET.parse(f).getroot()
        for pic in root.findall(".//xdr:pic", _NS):
            cNvPr = pic.find(".//xdr:cNvPr", _NS)
            if cNvPr is None:
                continue
            name = cNvPr.get("name")
            if not name or not name.startswith("ID_"):
                continue
            blip = pic.find(".//a:blip", _NS)
            if blip is None:
                continue
            rid = blip.get(f"{{{_NS['r']}}}embed")
            if rid:
                dispimg_to_rid[name] = rid

        # rId → 文件名
        rid_to_file: dict = {}
        if rels_path in names:
            with zf.open(rels_path) as f:
                rel_root = ET.parse(f).getroot()
            for rel in rel_root.findall(f".//{{{_REL_NS}}}Relationship"):
                rid = rel.get("Id")
                target = rel.get("Target")
                if rid and target:
                    rid_to_file[rid] = Path(target).name

        return {
            dispimg: rid_to_file[rid]
            for dispimg, rid in dispimg_to_rid.items()
            if rid in rid_to_file
        }