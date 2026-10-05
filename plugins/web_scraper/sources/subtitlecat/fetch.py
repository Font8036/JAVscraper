"""字幕猫爬取逻辑。"""

from __future__ import annotations

import asyncio
import logging
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional
from urllib.parse import quote, urljoin

from playwright.async_api import async_playwright
from openpyxl import load_workbook

from .config import SubtitleCatConfig
from .parse import (
    SearchResult,
    extract_result_count,
    parse_title,
    pick_best,
)

logger = logging.getLogger(__name__)

LogFn = Callable[[str], None]
ProgressFn = Callable[[int, str, dict, str], None]


# ============================================================
# 读 Excel
# ============================================================
def load_targets(excel_path: str | Path) -> list[str]:
    path = Path(excel_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel 不存在：{path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if not wb.worksheets:
            return []
        ws = wb.worksheets[0]
        targets: list[str] = []
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            v = row[0] if row else None
            if v is None:
                continue
            s = str(v).strip()
            if s and s.lower() not in ("nan", "null", "none"):
                targets.append(s)
        return targets
    finally:
        wb.close()


# ============================================================
# 主流程
# ============================================================
async def scrape_subtitlecat(
    config: SubtitleCatConfig,
    targets: list[str],
    extractor,
    *,
    stop_event=None,
    on_log: Optional[LogFn] = None,
    on_progress: Optional[ProgressFn] = None,
) -> list[dict]:
    log = on_log or (lambda _s: None)
    if not targets:
        return []

    records: list[dict] = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=not config.show_browser,
            channel=config.browser_channel or None,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )

        # 加载登录状态
        state_file = (config.state_file or "").strip()
        context_kwargs: dict = {"accept_downloads": True}
        if state_file and Path(state_file).exists():
            try:
                context_kwargs["storage_state"] = state_file
                log(f"已加载登录状态：{state_file}")
            except Exception as e:
                log(f"[警告] 加载登录状态失败：{e}")

        context = await browser.new_context(**context_kwargs)
        page = await context.new_page()
        page.set_default_timeout(config.page_timeout_ms)

        try:
            await page.goto(config.base_url, wait_until="load")
            log(f"已打开：{page.url}")

            # ============================================================
            # 第一轮：逐番号搜索
            # ============================================================
            pending: dict[int, dict] = {}   # idx -> {"href": ..., "deadline": ...}
            total = len(targets)

            for idx, keyword in enumerate(targets):
                if stop_event is not None and stop_event.is_set():
                    log("收到停止信号，中断爬取")
                    break

                log(f"\n===== [{idx+1}/{total}] {keyword} =====")
                if on_progress:
                    on_progress(idx, keyword, {}, "running")

                record = await _process_one(
                    page, keyword, config, extractor, log,
                )
                records.append(record)

                # 记录需要等待翻译的
                if record["status"] == "等待翻译":
                    pending[idx] = {
                        "href": record["_href"],
                        "deadline": time.monotonic() + max(
                            1, config.translate_timeout),
                    }

                if on_progress:
                    on_progress(idx, keyword, record, record["status"])

                # 返回搜索页
                try:
                    await page.goto(config.base_url, wait_until="load")
                except Exception:
                    pass

            # ============================================================
            # 后续轮：处理等待翻译
            # ============================================================
            while pending and not (stop_event is not None and stop_event.is_set()):
                log(f"\n--- 等待翻译重试，剩余 {len(pending)} 项，"
                    f"{config.round_delay:.1f}s 后继续 ---")
                await asyncio.sleep(max(0.1, config.round_delay))

                next_pending: dict[int, dict] = {}
                for idx, item in list(pending.items()):
                    keyword = targets[idx]
                    record = records[idx]
                    expired = time.monotonic() > item["deadline"]

                    # ---------- 先做一次尝试（不管是否已超时）----------
                    status = "下载失败"
                    try:
                        full_url = urljoin(config.base_url, item["href"])
                        await page.goto(full_url, wait_until="load")
                        status = await _try_download_or_translate(
                            page, keyword, config, log,
                        )
                    except Exception as e:
                        log(f"{keyword}: 重试异常：{e}")

                    # ---------- 按结果决定下一步 ----------
                    if status == "下载成功":
                        # 超时与否都算成功
                        if expired:
                            log(f"{keyword}: 超时后本次尝试成功下载")
                        record["status"] = _finalize_success(record, keyword)
                        if on_progress:
                            on_progress(idx, keyword, record, record["status"])

                    elif status == "等待翻译":
                        # 这次还是只有翻译按钮：这时才判超时
                        if expired:
                            record["status"] = "下载失败（翻译超时）"
                            log(f"{keyword}: 等待翻译超时")
                            if on_progress:
                                on_progress(idx, keyword, record, record["status"])
                        else:
                            next_pending[idx] = item

                    else:
                        # 下载失败 / 其他 → 直接失败
                        record["status"] = status
                        if on_progress:
                            on_progress(idx, keyword, record, record["status"])

                pending = next_pending

            # 未完成的标记已停止
            for idx in pending:
                records[idx]["status"] = "已停止"

            return records

        finally:
            try:
                await context.close()
            except Exception:
                pass
            try:
                await browser.close()
            except Exception:
                pass


# ============================================================
# 单番号处理
# ============================================================
async def _process_one(
    page,
    keyword: str,
    config: SubtitleCatConfig,
    extractor,
    log: LogFn,
) -> dict:
    """搜索一个关键词，选最佳结果，尝试下载。返回 record dict。"""
    record = _new_record(keyword, config)

    # 搜索
    try:
        await page.fill("#search", "")
        await page.fill("#search", keyword)
        await page.wait_for_selector("#button-addon2", timeout=10000)
        await page.click("#button-addon2")
    except Exception as e:
        log(f"{keyword}: 搜索提交失败：{e}")
        record["status"] = "下载失败"
        return record

    # 等待结果 / 空结果
    try:
        await page.wait_for_selector(
            ".nores__title, .sub-table tbody tr",
            state="visible",
            timeout=config.page_timeout_ms,
        )
    except Exception:
        log(f"{keyword}: 等待结果超时")
        record["status"] = "下载失败"
        return record

    # 无结果
    if await page.query_selector(".nores__title"):
        log(f"{keyword}: 无结果")
        record["status"] = "未搜到"
        return record

    # 结果数量
    try:
        sec_title = await page.inner_text(".sec-title")
        record["search_count"] = extract_result_count(sec_title)
    except Exception:
        record["search_count"] = 0

    # 解析
    results = await _parse_search_results(page)
    if not results:
        record["status"] = "未搜到"
        return record

    log(f"{keyword}: 共 {len(results)} 个结果")

    best = pick_best(results, keyword, extractor)
    if best is None:
        record["status"] = "未搜到"
        return record

    # 填表
    record["title"] = best.title_clean
    record["source_lang"] = best.source_lang
    record["downloads"] = best.downloads
    record["_href"] = best.href

    code = extractor.extract(best.title_clean)
    record["actual_code"] = code or ""

    star = "👍" if best.has_thumbs_up else "  "
    log(f"{keyword}: 选中 {best.title_clean} {star} "
        f"下载:{best.downloads} 语言:{best.source_lang}")

    # 进详情页
    full_url = urljoin(config.base_url, best.href)
    try:
        await page.goto(full_url, wait_until="load")
    except Exception as e:
        log(f"{keyword}: 进入详情页失败：{e}")
        record["status"] = "下载失败"
        return record

    # 尝试下载或点翻译
    status = await _try_download_or_translate(page, keyword, config, log)

    if status == "下载成功":
        record["status"] = _finalize_success(record, keyword)
    else:
        record["status"] = status

    return record


def _new_record(keyword: str, config: SubtitleCatConfig) -> dict:
    search_url = (
        f"{config.base_url.rstrip('/')}/index.php?search={quote(keyword)}"
    )
    return {
        "target": keyword,
        "search_count": 0,
        "actual_code": "",
        "title": "",
        "source_lang": "",
        "downloads": 0,
        "status": "",
        "_href": "",
        "_search_url": search_url,
    }


def _finalize_success(record: dict, keyword: str) -> str:
    """下载成功后，判断番号是否一致。"""
    from .parse import is_matched
    code = record.get("actual_code", "")
    if code and not is_matched(keyword, code):
        return "下载成功（番号不一致）"
    return "下载成功"


# ============================================================
# 解析搜索结果页
# ============================================================
async def _parse_search_results(page) -> list[SearchResult]:
    rows = await page.query_selector_all(".sub-table tbody tr")
    results: list[SearchResult] = []

    for row in rows:
        # 标题 / 链接
        link_el = await row.query_selector("td:nth-of-type(1) a")
        if not link_el:
            continue
        title_raw = (await link_el.inner_text()).strip()
        href = await link_el.get_attribute("href") or ""
        if not href:
            continue

        clean, lang_raw, lang_cn, is_cn = parse_title(title_raw)

        # 👍
        stars_el = await row.query_selector(".sub-table__stars")
        star_text = (await stars_el.inner_text()).strip() if stars_el else ""
        has_up = "👍" in star_text

        # 指标
        metric_els = await row.query_selector_all(".sub-table__metric-value")
        downloads = (
            await _parse_metric(metric_els[1])
            if len(metric_els) >= 2 else 0
        )
        languages = (
            await _parse_metric(metric_els[2])
            if len(metric_els) >= 3 else 0
        )

        results.append(SearchResult(
            title=title_raw,
            title_clean=clean,
            source_lang_raw=lang_raw,
            source_lang=lang_cn,
            translated_from_chinese=is_cn,
            has_thumbs_up=has_up,
            downloads=downloads,
            languages=languages,
            href=href,
        ))

    return results


async def _parse_metric(el) -> int:
    try:
        text = (await el.inner_text()).strip()
    except Exception:
        return 0
    import re
    m = re.search(r"(\d+)", text)
    return int(m.group(1)) if m else 0


# ============================================================
# 下载 / 翻译
# ============================================================
async def _try_download_or_translate(
    page,
    keyword: str,
    config: SubtitleCatConfig,
    log: LogFn,
) -> str:
    """检查当前页面，有下载按钮就下载，没下载有翻译按钮就点翻译。

    返回状态字符串："下载成功" / "等待翻译" / "下载失败"。
    """
    # 1. 有 #download_zh-CN
    try:
        dl = await page.query_selector("#download_zh-CN")
        if dl is not None and await dl.is_visible():
            return await _do_download(page, keyword, config, log)
    except Exception:
        pass

    # 2. 有 #zh-CN
    try:
        zh = await page.query_selector("#zh-CN")
        if zh is not None and await zh.is_visible():
            try:
                await zh.click()
                log(f"{keyword}: 已点击翻译按钮，等待生成…")
                return "等待翻译"
            except Exception as e:
                log(f"{keyword}: 点击翻译失败：{e}")
                return "下载失败"
    except Exception:
        pass

    return "下载失败"


async def _do_download(
    page,
    keyword: str,
    config: SubtitleCatConfig,
    log: LogFn,
) -> str:
    try:
        async with page.expect_download(
            timeout=config.page_timeout_ms,
        ) as dl_info:
            await page.click("#download_zh-CN")
        download = await dl_info.value

        filename = download.suggested_filename or f"{keyword}.srt"
        save_path = Path(config.target_dir) / filename
        save_path.parent.mkdir(parents=True, exist_ok=True)
        await download.save_as(str(save_path))
        log(f"{keyword}: ✓ 已保存 {save_path.name}")
        return "下载成功"
    except Exception as e:
        log(f"{keyword}: 下载失败：{e}")
        return "下载失败"